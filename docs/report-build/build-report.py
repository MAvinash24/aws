"""Build the AWS report from a working copy of the user's formatting reference."""
import argparse
import copy
import hashlib
import json
import zipfile
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table

ROOT=Path(__file__).resolve().parents[2]
ASSETS=ROOT/'docs/report-assets'

def field(paragraph,instruction):
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    code=OxmlElement('w:instrText');code.set(qn('xml:space'),'preserve');code.text=' '+instruction+' '
    separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    for element in (begin,code,separate,end):paragraph.add_run()._r.append(element)

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--reference',type=Path,required=True)
    parser.add_argument('--state',type=Path,default=ROOT/'reports/live-deployment.json')
    parser.add_argument('--output',type=Path,default=ROOT/'deliverables/AWS_DevSecOps_Capstone_Report_Final.docx')
    args=parser.parse_args()
    source_hash=hashlib.sha256(args.reference.read_bytes()).hexdigest()
    if source_hash!='b07f6a7f7f98ddad85524635ea71b51c876658f6b9f32915fc9e99ed8ee00388':raise RuntimeError('Formatting reference changed; inspect before rebuilding')
    state=json.loads(args.state.read_text(encoding='utf-8'))
    replacements={'RUN_ID':str(state['workflow']['id']),'SOURCE':state['release']['source'],
      'DIGEST':state['release']['image'].split('@')[1],'TASK_REVISION':state['release']['taskDefinition'].rsplit(':',1)[1]}
    def replace(text):
        for key,value in replacements.items():text=text.replace('{{'+key+'}}',value)
        return text
    doc=Document(args.reference)
    prototype=copy.deepcopy(doc.tables[4]._tbl)
    body=doc._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'):body.remove(child)
    for rel_id,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image'):del doc.part.rels[rel_id]
    for name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3','TOC Heading','Caption'):
        doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Title'].font.name='Aptos Display';doc.styles['Title'].font.size=Pt(32)
    doc.styles['Normal'].font.name='Aptos';doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    doc.styles['Heading 1'].paragraph_format.page_break_before=True
    for name in ('Title','Subtitle','TOC Heading','Heading 1','Heading 2','Heading 3'):
        pf=doc.styles[name].paragraph_format;pf.keep_with_next=True
        borders=doc.styles[name]._element.xpath('./w:pPr/w:pBdr')
        for border in borders:border.getparent().remove(border)
    section=doc.sections[0]
    for header in (section.header,section.first_page_header,section.even_page_header):
        for p in header.paragraphs:
            for r in p.runs:
                if 'SkillSim' in r.text:r.text=r.text.replace('SkillSim','AWS DevSecOps')
                r.font.color.rgb=RGBColor(0,0,0)
    table_count=0;figure_count=0
    def paragraph(text,style=None):
        p=doc.add_paragraph(replace(text),style=style)
        p.paragraph_format.widow_control=True
        p.paragraph_format.keep_together=True
        if not style:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        return p
    def caption(text):
        p=paragraph(text,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
        return p
    def add_table(block,cover=False):
        nonlocal table_count
        headers=block['headers'];rows=block['rows'];columns=len(headers)
        element=copy.deepcopy(prototype)
        for child in list(element):
            if child.tag in (qn('w:tr'),qn('w:tblGrid')):element.remove(child)
        grid=OxmlElement('w:tblGrid')
        widths=block.get('widths') or [7.06/columns]*columns
        for width in widths:
            c=OxmlElement('w:gridCol');c.set(qn('w:w'),str(round(width*1440)));grid.append(c)
        element.insert(1,grid)
        body.insert(len(body)-1,element)
        t=Table(element,doc._body);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        props=t._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for name in ('top','left','bottom','right','insideH','insideV'):
            border=OxmlElement('w:'+name);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
        props.append(borders)
        for row_num,values in enumerate([headers]+rows):
            row=t.add_row()
            trPr=row._tr.get_or_add_trPr()
            no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
            if row_num==0:
                repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
            for i,(cell,value) in enumerate(zip(row.cells,values)):
                cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17365D' if row_num==0 else ('F2F6FA' if row_num%2==0 else 'FFFFFF'));cp.append(shade)
                margins=OxmlElement('w:tcMar')
                for side,num in (('top',70),('bottom',70),('left',90),('right',90)):
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(num));e.set(qn('w:type'),'dxa');margins.append(e)
                cp.append(margins)
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
                if i==0 and widths[i]<1.2:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run(replace(str(value)));r.font.name='Aptos';r.font.size=Pt(9.2 if columns<=3 else 8.6)
                r.font.bold=(row_num==0);r.font.color.rgb=RGBColor.from_string('FFFFFF' if row_num==0 else '172033')
        if not cover:
            table_count+=1
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=Pt(3);p.paragraph_format.space_before=Pt(0)
        return t
    def add_figure(block):
        nonlocal figure_count
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(0)
        r=p.add_run();image=r.add_picture(str(ASSETS/block['file']),width=Inches(6.8))
        desc=image._inline.docPr;desc.set('descr',block['caption']);desc.set('title',block['file'])
        figure_count+=1
        p=doc.add_paragraph(style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
        p.add_run('Figure ');field(p,'SEQ Figure \\* ARABIC');p.add_run('  '+replace(block['caption']))
    p=paragraph('SECURE SOFTWARE ENGINEERING CAPSTONE REPORT');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(36);p.paragraph_format.space_after=Pt(15)
    for r in p.runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('5B6573')
    p=paragraph('AWS DevSecOps','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=paragraph('Secure Container Delivery and Runtime Monitoring','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=paragraph('GitHub Actions  AWS ECR and ECS  Falco  Local Docker');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(25)
    add_table({'headers':['Document Control','Details'],'rows':[
      ['Author','Marisetti Avinash'],['Student ID','CH.SC.U4CYS23024'],['Institution','Amrita Vishwa Vidyapeetham Chennai'],['Supervisor','S.Udhaya Kumar'],['Subject','Secure Software Engineering'],['Project','AWS DevSecOps secure delivery demonstration'],['Repository','https://github.com/MAvinash24/aws'],['Region','Mumbai ap-south-1'],['Report date','6 October 2026'],['Verified release',replacements['RUN_ID']]],'widths':[1.8,5.26]},cover=True)
    paragraph('The implemented pipeline has completed a real signed deployment to AWS. The application is healthy in ECS and in the independent local Docker demonstration. This report records the engineering design, final changes, verification evidence and remaining production limits.').paragraph_format.space_before=Pt(15)
    doc.add_page_break()
    paragraph('Table of Contents','TOC Heading')
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    paragraph('List of Figures','TOC Heading')
    field(doc.add_paragraph(), 'TOC \\c "Figure" \\h \\z')
    content=json.loads(Path(__file__).with_name('report-content.json').read_text(encoding='utf-8'))
    for chapter in content:
        paragraph(chapter['title'],'Heading 1');paragraph(chapter['summary'])
        for s in chapter['sections']:
            paragraph(s['title'],'Heading 2')
            for block in s['blocks']:
                if block['type']=='paragraph':paragraph(block['text'])
                elif block['type']=='table':
                    p=caption('Table '+str(table_count+1)+'  '+block['title']);p.paragraph_format.keep_with_next=True
                    add_table(block)
                elif block['type']=='figure':add_figure(block)
                elif block['type']=='code':
                    for line in block['text'].splitlines():
                        p=paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(1)
                        for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
    doc.core_properties.title='AWS DevSecOps Secure Software Engineering Capstone Report'
    doc.core_properties.subject='Verified secure container delivery and runtime monitoring'
    doc.core_properties.author='Marisetti Avinash';doc.core_properties.last_modified_by='Marisetti Avinash'
    doc.core_properties.comments='';doc.core_properties.keywords='AWS, DevSecOps, GitHub Actions, ECS, Falco'
    settings=doc.settings._element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    args.output.parent.mkdir(exist_ok=True)
    doc.save(args.output)
    # Preserve unrelated opaque design package parts byte for byte.
    preserve=['word/theme/theme1.xml','word/numbering.xml','word/fontTable.xml']
    with zipfile.ZipFile(args.reference) as source,zipfile.ZipFile(args.output) as generated:
        parts={n:generated.read(n) for n in generated.namelist()}
        for name in preserve:
            if name in source.namelist():parts[name]=source.read(name)
    with zipfile.ZipFile(args.output,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,data in parts.items():archive.writestr(name,data)
    if hashlib.sha256(args.reference.read_bytes()).hexdigest()!=source_hash:raise RuntimeError('Reference was modified')
    print('Created',args.output,'chapters',len(content),'tables',table_count+1,'figures',figure_count)

if __name__=='__main__':main()
